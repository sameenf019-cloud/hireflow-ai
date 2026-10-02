"""Document parsing + Qdrant/FastEmbed hybrid retrieval (dense + BM25 sparse, fused with RRF).

Heavy dependencies (fitz, python-docx, qdrant-client, fastembed) are imported lazily so that
importing this module is cheap and the pure helpers (chunking, email extraction) are testable.

NOTE: with the embedded local Qdrant (no QDRANT_URL) only ONE process can open the storage
folder at a time. Run a single uvicorn worker, or set QDRANT_URL for a server/cloud instance.
"""
import re
import threading
import uuid
from functools import lru_cache
from pathlib import Path
from typing import Any, Optional

from config import settings

DENSE_VECTOR_NAME = "dense"
SPARSE_VECTOR_NAME = "bm25"
_POINT_NAMESPACE = uuid.UUID("6f1c1d3e-5b0a-4c63-9d57-0b6c2a1f4e11")
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
_collection_lock = threading.Lock()

# ===================================================================== parsing


def extract_text(path: str | Path) -> str:
    """Extract plain text from a PDF (PyMuPDF), DOCX (python-docx) or TXT file."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import fitz  # PyMuPDF

        with fitz.open(path) as doc:
            text = "\n".join(page.get_text("text") for page in doc)
    elif suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type '{suffix}'. Use PDF, DOCX or TXT.")

    text = re.sub(r"[ \t]+", " ", text.replace("\r", "")).strip()
    if not text:
        raise ValueError(f"No extractable text in {path.name} (scanned PDF? OCR is not supported).")
    return text


def extract_email(text: str) -> Optional[str]:
    match = _EMAIL_RE.search(text)
    return match.group(0).lower() if match else None


def guess_name(text: str) -> Optional[str]:
    """Resumes almost always open with the candidate's name: take the first short, letters-only line."""
    for line in text.splitlines()[:8]:
        line = line.strip()
        if 2 <= len(line) <= 50 and "@" not in line and re.fullmatch(r"[A-Za-z][A-Za-z .'\-]+", line):
            if 1 < len(line.split()) <= 5:
                return line.title() if line.isupper() else line
    return None


# ==================================================================== chunking


def _window(text: str, size: int, overlap: int) -> list[str]:
    """Split one over-long line into overlapping windows, snapping to word boundaries."""
    out, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            space = text.rfind(" ", start, end)
            if space > start + size // 2:
                end = space
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [w for w in out if w]


def chunk_text(text: str, size: Optional[int] = None, overlap: Optional[int] = None) -> list[str]:
    """Line-aware chunking: lines are packed up to `size` chars, with `overlap` chars of
    context carried into the next chunk. Keeps section headings next to their bullet lists."""
    size = size or settings.chunk_size
    overlap = settings.chunk_overlap if overlap is None else overlap
    pieces: list[str] = []
    for line in (ln.strip() for ln in text.splitlines()):
        if not line:
            continue
        pieces.extend([line] if len(line) <= size else _window(line, size, overlap))

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + 1 + len(piece) > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            if " " in tail:
                tail = tail[tail.find(" ") + 1:]  # start the overlap on a word boundary
            current = f"{tail}\n{piece}" if tail else piece
        else:
            current = f"{current}\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks


# ============================================================ Qdrant / FastEmbed


@lru_cache(maxsize=1)
def _dense_model():
    from fastembed import TextEmbedding

    return TextEmbedding(model_name=settings.dense_model)


@lru_cache(maxsize=1)
def _sparse_model():
    from fastembed import SparseTextEmbedding

    return SparseTextEmbedding(model_name=settings.sparse_model)


@lru_cache(maxsize=1)
def get_client():
    from qdrant_client import QdrantClient

    if settings.qdrant_url:
        return QdrantClient(url=settings.qdrant_url, api_key=settings.qdrant_api_key or None)
    settings.ensure_dirs()
    return QdrantClient(path=str(settings.qdrant_local_path))


def ensure_collection() -> None:
    from fastembed import TextEmbedding
    from qdrant_client import models

    client = get_client()
    with _collection_lock:
        if client.collection_exists(settings.qdrant_collection):
            return
        client.create_collection(
            collection_name=settings.qdrant_collection,
            vectors_config={
                DENSE_VECTOR_NAME: models.VectorParams(
                    size=TextEmbedding.get_embedding_size(settings.dense_model),
                    distance=models.Distance.COSINE,
                )
            },
            sparse_vectors_config={
                SPARSE_VECTOR_NAME: models.SparseVectorParams(modifier=models.Modifier.IDF)
            },
        )
        for field in ("candidate_id", "job_id"):
            try:
                client.create_payload_index(
                    settings.qdrant_collection, field, models.PayloadSchemaType.INTEGER
                )
            except Exception:  # local mode ignores/does not need payload indexes
                pass


def _to_sparse(embedding) -> Any:
    from qdrant_client import models

    return models.SparseVector(indices=embedding.indices.tolist(), values=embedding.values.tolist())


def delete_candidate_vectors(candidate_id: int) -> None:
    from qdrant_client import models

    ensure_collection()
    get_client().delete(
        collection_name=settings.qdrant_collection,
        points_selector=models.FilterSelector(
            filter=models.Filter(
                must=[models.FieldCondition(key="candidate_id", match=models.MatchValue(value=candidate_id))]
            )
        ),
    )


def index_resume(
    candidate_id: int,
    job_id: int,
    email: str,
    text: str,
    filename: Optional[str] = None,
) -> int:
    """Chunk -> embed (dense + BM25) -> upsert. Re-indexing a candidate replaces old chunks.
    Returns the number of chunks stored."""
    from qdrant_client import models

    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("Resume produced no chunks.")
    ensure_collection()
    delete_candidate_vectors(candidate_id)

    dense_vecs = [v.tolist() for v in _dense_model().embed(chunks)]
    sparse_vecs = [_to_sparse(s) for s in _sparse_model().embed(chunks)]

    points = [
        models.PointStruct(
            id=str(uuid.uuid5(_POINT_NAMESPACE, f"{candidate_id}:{i}")),
            vector={DENSE_VECTOR_NAME: dense_vecs[i], SPARSE_VECTOR_NAME: sparse_vecs[i]},
            payload={
                "candidate_id": candidate_id,
                "job_id": job_id,
                "email": email,
                "filename": filename,
                "chunk_index": i,
                "text": chunk,
            },
        )
        for i, chunk in enumerate(chunks)
    ]
    get_client().upsert(collection_name=settings.qdrant_collection, points=points)
    return len(points)


def hybrid_search(
    query: str,
    candidate_id: Optional[int] = None,
    job_id: Optional[int] = None,
    limit: int = 5,
) -> list[dict[str, Any]]:
    """Dense (semantic) + BM25 (exact keyword) retrieval fused with Reciprocal Rank Fusion."""
    from qdrant_client import models

    query = query.strip()
    if not query:
        return []
    ensure_collection()

    must = []
    if candidate_id is not None:
        must.append(models.FieldCondition(key="candidate_id", match=models.MatchValue(value=candidate_id)))
    if job_id is not None:
        must.append(models.FieldCondition(key="job_id", match=models.MatchValue(value=job_id)))
    query_filter = models.Filter(must=must) if must else None

    dense_q = next(iter(_dense_model().query_embed(query))).tolist()
    sparse_q = _to_sparse(next(iter(_sparse_model().query_embed(query))))
    prefetch_limit = max(limit * 4, 20)

    result = get_client().query_points(
        collection_name=settings.qdrant_collection,
        prefetch=[
            models.Prefetch(query=dense_q, using=DENSE_VECTOR_NAME, filter=query_filter, limit=prefetch_limit),
            models.Prefetch(query=sparse_q, using=SPARSE_VECTOR_NAME, filter=query_filter, limit=prefetch_limit),
        ],
        query=models.FusionQuery(fusion=models.Fusion.RRF),
        query_filter=query_filter,
        limit=limit,
        with_payload=True,
    )
    return [
        {
            "score": float(p.score),
            "candidate_id": p.payload["candidate_id"],
            "job_id": p.payload["job_id"],
            "email": p.payload.get("email"),
            "filename": p.payload.get("filename"),
            "chunk_index": p.payload["chunk_index"],
            "text": p.payload["text"],
        }
        for p in result.points
    ]


def _requirement_lines(jd_text: str, max_lines: int = 8) -> list[str]:
    """Bullet / numbered lines from the JD: each becomes its own query so exact keyword
    requirements (BM25) are checked individually, not diluted by the whole JD."""
    lines = []
    for raw in jd_text.splitlines():
        line = raw.strip()
        if re.match(r"^([-•*●▪·]|\d+[.)])\s+\S", line) and 8 <= len(line) <= 200:
            lines.append(re.sub(r"^([-•*●▪·]|\d+[.)])\s+", "", line))
        if len(lines) >= max_lines:
            break
    return lines


def build_screening_context(candidate_id: int, jd_text: str, limit: int = 8) -> str:
    """Evidence passages from one resume for the Screening Agent: one semantic query with the
    whole JD plus one query per requirement line, merged and de-duplicated by chunk."""
    queries = [jd_text[:1500]] + _requirement_lines(jd_text)
    best: dict[int, dict[str, Any]] = {}
    for q in queries:
        for hit in hybrid_search(q, candidate_id=candidate_id, limit=4):
            prev = best.get(hit["chunk_index"])
            if prev is None or hit["score"] > prev["score"]:
                best[hit["chunk_index"]] = hit
    top = sorted(best.values(), key=lambda h: h["score"], reverse=True)[:limit]
    top.sort(key=lambda h: h["chunk_index"])  # restore resume order for readability
    if not top:
        return "No resume passages found for this candidate."
    return "\n\n".join(f"[Resume passage {h['chunk_index'] + 1}]\n{h['text']}" for h in top)


# ================================================================ ingestion (Stage 1)


def ingest_job_description(path: str | Path, title: Optional[str] = None) -> int:
    """Parse a JD file and store it in SQLite. Returns job_id."""
    import state_manager as db

    path = Path(path)
    text = extract_text(path)
    return db.create_job(title=title or path.stem.replace("_", " ").strip(), jd_text=text, filename=path.name)


def ingest_resume(job_id: int, path: str | Path, email_override: Optional[str] = None) -> dict[str, Any]:
    """Stage 1 for one resume: parse -> SQLite row (PENDING_SCREENING) -> chunk/embed -> Qdrant.

    If no email can be found in the resume and no override is given, a clearly fake placeholder
    address is used and `email_found` is False so the API can ask the recruiter to supply one."""
    import state_manager as db

    path = Path(path)
    text = extract_text(path)
    email = (email_override or extract_email(text) or "").strip().lower()
    email_found = bool(email)
    if not email_found:
        slug = re.sub(r"[^a-z0-9]+", "-", path.stem.lower()).strip("-") or "candidate"
        email = f"{slug}@unknown.invalid"

    candidate_id = db.create_candidate(
        job_id=job_id, email=email, resume_text=text, name=guess_name(text), resume_filename=path.name
    )
    chunks = index_resume(candidate_id, job_id, email, text, path.name)
    db.log_event("Ingestion", f"Parsed and indexed {path.name} ({chunks} chunks)", candidate_id)
    return {
        "candidate_id": candidate_id,
        "email": email,
        "email_found": email_found,
        "name": guess_name(text),
        "chunks": chunks,
    }
